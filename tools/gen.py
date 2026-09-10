import openpyxl, json, collections
wb = openpyxl.load_workbook(r"C:\Users\Bo\Downloads\2026 NFL .xlsx", data_only=True)
ws = wb["2026 NFL"]
games=[]
for r in range(2, 359):
    a = ws.cell(r,1).value
    if isinstance(a,str) and a.upper().startswith("WEEK"):
        wk=int(a.split()[1])
        note = ws.cell(r,15).value
        g=[wk, ws.cell(r,2).value, ws.cell(r,3).value, ws.cell(r,8).value,
           ws.cell(r,9).value, ws.cell(r,10).value]
        if note: g.append(note)
        games.append(g)
assert len(games)==272, len(games)
# validate
teams=sorted({g[1] for g in games} | {g[2] for g in games})
assert len(teams)==32
for g in games:
    assert g[3] in (g[1],g[2]), g
    if not isinstance(g[4],(int,float)):
        print("NO SPREAD:",g); g[4]=None
# byes
byes={}
for t in teams:
    played={g[0] for g in games if g[1]==t or g[2]==t}
    b=sorted(set(range(1,19))-played)
    byes[t]=b
    assert len(b)==1, (t,b)
# home/away counts
for t in teams:
    h=sum(1 for g in games if g[2]==t); a=sum(1 for g in games if g[1]==t)
    assert h+a==17, (t,h,a)
print("byes:", {t:byes[t][0] for t in teams})
fav=collections.Counter(g[3] for g in games)
print("fav counts:", sorted(fav.items(), key=lambda x:-x[1]))
out={"games":games,"byes":{t:byes[t][0] for t in teams}}
json.dump(out, open("data2026.json","w"), separators=(",",":"))
print("written", len(json.dumps(out)))
